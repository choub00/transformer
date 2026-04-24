# -*- coding: utf-8 -*-
from docx import Document
import re

OUTPUT = r"D:/transformer/thesis/基于transformer的股票预测系统_润色版.docx"

doc = Document(OUTPUT)

fixes = [
    ("因此", "所以"),
    ("本课题", "本研究"),
]

total_changes = 0

for para in doc.paragraphs:
    for run in para.runs:
        if run.text:
            for old, new in fixes:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    total_changes += 1
                    print(f"Fixed '{old}' -> '{new}' in paragraph: {para.text[:60]}")

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if run.text:
                        for old, new in fixes:
                            if old in run.text:
                                run.text = run.text.replace(old, new)
                                total_changes += 1
                                print(f"Fixed '{old}' -> '{new}' in TABLE")

print(f"\nTotal fixes applied: {total_changes}")

doc.save(OUTPUT)
print(f"Saved to: {OUTPUT}")
