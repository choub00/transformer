from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


BASE_DIR = Path(r"D:\transformer\thesis")
REFERENCE_DOC = BASE_DIR / "参考论文-基于自然语言处理的银行智能语音助手的设计与实现-智能系统类.docx"
SOURCE_DOC = BASE_DIR / "基于Transformer的股票预测系统_论文初稿.docx"
OUTPUT_DOC = BASE_DIR / "基于Transformer的股票预测系统_参考格式修订版.docx"


RE_CHAPTER = re.compile(r"^(第[一二三四五六七八九十0-9]+章\s*.+|\d+\s+.+)$")
RE_SECTION = re.compile(r"^\d+\.\d+\s+.+$")
RE_SUBSECTION = re.compile(r"^\d+\.\d+\.\d+\s+.+$")
RE_FIGURE = re.compile(r"^(图|表|续表)\s*\d+[-—]\d+")
RE_KEYWORDS = re.compile(r"^(关键词|Key words)\s*[：:]")


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)


def set_run_font(run, font_name: str, size_pt: float, bold: bool | None = None) -> None:
    set_east_asia_font(run, font_name)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def para_text(para) -> str:
    return para.text.strip().replace("\u3000", " ")


def is_ascii_text(text: str) -> bool:
    return bool(text) and not re.search(r"[\u4e00-\u9fff]", text)


def looks_like_formula(text: str) -> bool:
    if not text:
        return False
    if re.search(r"[\u4e00-\u9fff]", text):
        return False
    formula_tokens = [
        "\\frac", "\\math", "\\hat", "\\bar", "\\tilde", "\\sum", "\\prod",
        "\\sqrt", "\\left", "\\right", "\\text", "\\mathrm", "\\mathbf",
        "\\begin", "\\end", "\\cdot", "\\top", "\\in", "\\cup", "\\cap",
        "^", "_", "=", "\\quad", "\\qquad",
    ]
    return any(token in text for token in formula_tokens) or text.startswith("\\")


def set_paragraph_common(para, align=None, first_indent=None, space_before=0, space_after=0, exact_line_pt=20) -> None:
    fmt = para.paragraph_format
    para.alignment = align
    fmt.first_line_indent = first_indent
    fmt.left_indent = None
    fmt.right_indent = None
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(exact_line_pt)


def format_runs_uniform(para, font_name: str, size_pt: float, bold: bool | None = None) -> None:
    for run in para.runs:
        if not run.text:
            continue
        set_run_font(run, font_name, size_pt, bold)


def get_abstract_start(doc: Document) -> int:
    for i, para in enumerate(doc.paragraphs):
        if para_text(para) == "摘 要":
            return i
    return 0


def format_cover(doc: Document) -> None:
    abstract_start = get_abstract_start(doc)
    nonempty = [p for p in doc.paragraphs[:abstract_start] if para_text(p)]
    if len(nonempty) < 6:
        return

    for idx, para in enumerate(nonempty):
        text = para_text(para)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt = para.paragraph_format
        fmt.first_line_indent = None
        fmt.left_indent = None
        fmt.right_indent = None
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY

        if idx == 0:
            fmt.space_before = Pt(90)
            fmt.space_after = Pt(18)
            fmt.line_spacing = Pt(36)
            format_runs_uniform(para, "黑体", 26, True)
        elif idx == 1:
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(12)
            fmt.line_spacing = Pt(24)
            format_runs_uniform(para, "宋体", 16, False)
        elif idx < len(nonempty) - 1:
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(6)
            fmt.line_spacing = Pt(22)
            format_runs_uniform(para, "宋体", 14, False)
        else:
            fmt.space_before = Pt(18)
            fmt.space_after = Pt(0)
            fmt.line_spacing = Pt(24)
            format_runs_uniform(para, "宋体", 16, False)


def format_body(doc: Document) -> None:
    abstract_start = get_abstract_start(doc)

    for idx, para in enumerate(doc.paragraphs):
        text = para_text(para)
        if not text:
            continue

        if idx < abstract_start:
            continue

        if text in {"摘 要", "Abstract", "目 录", "参考文献", "致谢"}:
            set_paragraph_common(para, WD_ALIGN_PARAGRAPH.CENTER, None, 18, 12, 20)
            font = "黑体" if re.search(r"[\u4e00-\u9fff]", text) else "Times New Roman"
            format_runs_uniform(para, font, 16, True)
            continue

        if RE_KEYWORDS.match(text):
            set_paragraph_common(para, WD_ALIGN_PARAGRAPH.LEFT, Cm(0.84), 0, 0, 20)
            for i, run in enumerate(para.runs):
                if not run.text:
                    continue
                if i == 0 and ("关键词" in run.text or "Key words" in run.text):
                    font = "黑体" if "关键词" in run.text else "Times New Roman"
                    set_run_font(run, font, 12, True)
                else:
                    font = "宋体" if re.search(r"[\u4e00-\u9fff]", run.text) else "Times New Roman"
                    set_run_font(run, font, 12, False)
            continue

        if RE_CHAPTER.match(text):
            set_paragraph_common(para, WD_ALIGN_PARAGRAPH.LEFT, None, 12, 6, 20)
            format_runs_uniform(para, "黑体", 15, True)
            continue

        if RE_SUBSECTION.match(text):
            set_paragraph_common(para, WD_ALIGN_PARAGRAPH.LEFT, None, 6, 3, 20)
            format_runs_uniform(para, "黑体", 13, False)
            continue

        if RE_SECTION.match(text):
            set_paragraph_common(para, WD_ALIGN_PARAGRAPH.LEFT, None, 9, 3, 20)
            format_runs_uniform(para, "黑体", 14, False)
            continue

        if RE_FIGURE.match(text) or text.startswith("[此处插入图") or text.startswith("[此处插入表"):
            set_paragraph_common(para, WD_ALIGN_PARAGRAPH.CENTER, None, 6, 6, 20)
            format_runs_uniform(para, "黑体", 10.5, False)
            continue

        if looks_like_formula(text):
            set_paragraph_common(para, WD_ALIGN_PARAGRAPH.CENTER, None, 6, 6, 20)
            format_runs_uniform(para, "Cambria Math", 11, False)
            continue

        if is_ascii_text(text):
            set_paragraph_common(para, WD_ALIGN_PARAGRAPH.JUSTIFY, None, 0, 0, 20)
            format_runs_uniform(para, "Times New Roman", 12, False)
            continue

        set_paragraph_common(para, WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(0.84), 0, 0, 20)
        for run in para.runs:
            if not run.text:
                continue
            font = "宋体" if re.search(r"[\u4e00-\u9fff]", run.text) else "Times New Roman"
            set_run_font(run, font, 12, False)


def copy_page_setup(reference: Document, target: Document) -> None:
    ref_sec = reference.sections[0]
    for sec in target.sections:
        sec.page_width = ref_sec.page_width
        sec.page_height = ref_sec.page_height
        sec.top_margin = ref_sec.top_margin
        sec.bottom_margin = ref_sec.bottom_margin
        sec.left_margin = ref_sec.left_margin
        sec.right_margin = ref_sec.right_margin
        sec.header_distance = ref_sec.header_distance
        sec.footer_distance = ref_sec.footer_distance


def main() -> None:
    ref_doc = Document(str(REFERENCE_DOC))
    src_doc = Document(str(SOURCE_DOC))

    copy_page_setup(ref_doc, src_doc)
    format_cover(src_doc)
    format_body(src_doc)
    src_doc.save(str(OUTPUT_DOC))
    print(f"saved: {OUTPUT_DOC}")


if __name__ == "__main__":
    main()
