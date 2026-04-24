from __future__ import annotations

import math
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
THESIS_DIR = ROOT / "thesis"
DOCX_PATH = THESIS_DIR / "22.15.docx"
FIG_DIR = THESIS_DIR / "generated_figures"


def find_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    font_candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc") if bold else Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for path in font_candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


FONT_TITLE = find_font(42, True)
FONT_SUBTITLE = find_font(28, True)
FONT_BODY = find_font(24)
FONT_SMALL = find_font(20)
FONT_TINY = find_font(18)


def text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        trial = current + char
        if text_size(draw, trial, font)[0] <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def rounded_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    text: str,
    fill: str,
    outline: str = "#2f4a60",
    font: ImageFont.ImageFont = FONT_BODY,
    text_fill: str = "#0f1720",
    radius: int = 22,
    width: int = 3,
):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    x1, y1, x2, y2 = xy
    lines = wrap_text(draw, text, font, x2 - x1 - 36)
    line_h = text_size(draw, "国", font)[1] + 8
    total_h = line_h * len(lines)
    y = y1 + (y2 - y1 - total_h) // 2
    for line in lines:
        tw, th = text_size(draw, line, font)
        draw.text((x1 + (x2 - x1 - tw) // 2, y), line, font=font, fill=text_fill)
        y += line_h


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    fill: str = "#315b7d",
    width: int = 5,
):
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    spread = 0.45
    p1 = (
        end[0] - length * math.cos(angle - spread),
        end[1] - length * math.sin(angle - spread),
    )
    p2 = (
        end[0] - length * math.cos(angle + spread),
        end[1] - length * math.sin(angle + spread),
    )
    draw.polygon([end, p1, p2], fill=fill)


def canvas(title: str, subtitle: str = "") -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (1600, 900), "#f8fbfd")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 1600, 110), fill="#e8f2f7")
    draw.text((70, 35), title, font=FONT_TITLE, fill="#19324a")
    if subtitle:
        draw.text((70, 88), subtitle, font=FONT_SMALL, fill="#536b7c")
    return img, draw


def save(img: Image.Image, filename: str) -> Path:
    FIG_DIR.mkdir(exist_ok=True)
    path = FIG_DIR / filename
    img.save(path, "PNG")
    return path


def fig_3_1() -> Path:
    img, draw = canvas("系统四层总体架构", "数据层、模型层、服务层与展示层协同支撑股票预测系统")
    layers = [
        ("数据层", "历史行情\n股票池维护\n特征工程\n滚动标准化", "#dceefb"),
        ("模型层", "AlphaTransformer v3\n训练与验证\n模型注册\n推理输出", "#e6f4ea"),
        ("服务层", "FastAPI\n预测接口\n回测接口\n账户与交易接口", "#fff4d6"),
        ("展示层", "Vue 3\nDashboard\n分析页\n交易页", "#f3e8ff"),
    ]
    y = 180
    for i, (name, desc, fill) in enumerate(layers):
        x = 105 + i * 375
        rounded_box(draw, (x, y, x + 280, y + 110), name, fill, font=FONT_SUBTITLE)
        rounded_box(draw, (x, y + 155, x + 280, y + 430), desc, "#ffffff", font=FONT_BODY)
        if i < len(layers) - 1:
            arrow(draw, (x + 300, y + 310), (x + 365, y + 310))
    draw.text((175, 735), "输入：多资产价量数据", font=FONT_BODY, fill="#1f5673")
    draw.text((605, 735), "输出：排序信号、回测指标、可视化结果", font=FONT_BODY, fill="#1f5673")
    return save(img, "fig_3_1_system_architecture.png")


def fig_3_2() -> Path:
    img, draw = canvas("核心业务流程", "从原始行情到前端展示的闭环研究链路")
    steps = [
        "数据清洗\n特征衍生",
        "滚动标准化\n样本构建",
        "模型训练\nWalk-Forward",
        "回测评估\n交易成本",
        "API 封装\n结构化响应",
        "前端展示\n决策交互",
    ]
    xs = [80, 330, 580, 830, 1080, 1330]
    for i, step in enumerate(steps):
        rounded_box(draw, (xs[i], 280, xs[i] + 190, 440), step, "#ffffff", font=FONT_BODY)
        draw.ellipse((xs[i] + 70, 190, xs[i] + 120, 240), fill="#2e78a6")
        draw.text((xs[i] + 84, 199), str(i + 1), font=FONT_BODY, fill="#ffffff")
        if i < len(steps) - 1:
            arrow(draw, (xs[i] + 205, 360), (xs[i + 1] - 15, 360))
    rounded_box(draw, (340, 590, 1260, 695), "统一数据口径 + 统一评估协议 + 统一接口输出", "#e8f2f7", font=FONT_SUBTITLE)
    return save(img, "fig_3_2_business_flow.png")


def fig_4_1() -> Path:
    img, draw = canvas("AlphaTransformer v3 总体架构", "论文方案示意图：状态感知、时间记忆与资产联动分工建模")
    boxes = [
        ("多资产历史窗口\n[B,A,T,F]", "#dceefb"),
        ("Embedding\n统一隐空间", "#e6f4ea"),
        ("Regime-Aware\n市场状态概率", "#fff4d6"),
        ("Mamba-2\n时间记忆递推", "#f3e8ff"),
        ("iTransformer\n资产维联动", "#fde2e2"),
        ("Output Head\n排序分数", "#e8f2f7"),
    ]
    xs = [55, 310, 565, 820, 1075, 1330]
    for i, (label, fill) in enumerate(boxes):
        rounded_box(draw, (xs[i], 300, xs[i] + 205, 455), label, fill, font=FONT_BODY)
        if i < len(boxes) - 1:
            arrow(draw, (xs[i] + 218, 378), (xs[i + 1] - 15, 378))
    rounded_box(draw, (345, 600, 1255, 715), "CombinedQuantLoss：预测误差 + 横截面排序 + 换手/交易成本约束", "#ffffff", font=FONT_BODY)
    arrow(draw, (1435, 470), (1100, 600), fill="#7c4d9a")
    return save(img, "fig_4_1_v3_architecture.png")


def fig_4_2() -> Path:
    img, draw = canvas("Regime-Aware 状态感知机制", "用重建误差与池化特征生成状态概率，并动态调节归一化参数")
    rounded_box(draw, (90, 280, 330, 450), "输入特征窗口\n价量与技术因子", "#dceefb", font=FONT_BODY)
    rounded_box(draw, (430, 210, 700, 360), "Autoencoder\n重建输入", "#e6f4ea", font=FONT_BODY)
    rounded_box(draw, (430, 440, 700, 590), "时序池化特征\n趋势/波动摘要", "#e6f4ea", font=FONT_BODY)
    rounded_box(draw, (805, 320, 1085, 480), "状态概率\n平稳 / 波动 / 冲击", "#fff4d6", font=FONT_BODY)
    rounded_box(draw, (1190, 270, 1500, 530), "Adaptive Layer Norm\nγ(s)、β(s)\n动态缩放与平移", "#f3e8ff", font=FONT_BODY)
    arrow(draw, (330, 365), (430, 285))
    arrow(draw, (330, 365), (430, 515))
    arrow(draw, (700, 285), (805, 355))
    arrow(draw, (700, 515), (805, 445))
    arrow(draw, (1085, 400), (1190, 400))
    draw.text((510, 380), "重建误差衡量偏离常态程度", font=FONT_SMALL, fill="#475569")
    return save(img, "fig_4_2_regime_aware.png")


def fig_4_3() -> Path:
    img, draw = canvas("Mamba-2 时间编码与记忆衰减", "以递推状态更新保留有效信号、衰减陈旧噪声")
    y = 430
    for i in range(6):
        x = 120 + i * 240
        draw.ellipse((x, y - 45, x + 90, y + 45), fill="#ffffff", outline="#315b7d", width=4)
        draw.text((x + 22, y - 16), f"t{i+1}", font=FONT_BODY, fill="#1f3d56")
        if i < 5:
            arrow(draw, (x + 95, y), (x + 225, y))
        draw.rectangle((x - 15, y + 110, x + 105, y + 145), fill=["#a7f3d0", "#bbf7d0", "#d9f99d", "#fde68a", "#fed7aa", "#fecaca"][i])
    rounded_box(draw, (280, 185, 1310, 280), "隐状态 h_t 持续更新：近期趋势、成交放量、波动抬升获得更高权重", "#e8f2f7", font=FONT_BODY)
    rounded_box(draw, (420, 650, 1180, 760), "远期无效扰动逐步衰减，避免标准注意力把所有历史点等强度比较", "#ffffff", font=FONT_BODY)
    return save(img, "fig_4_3_mamba_memory.png")


def fig_4_4() -> Path:
    img, draw = canvas("时间资产维编码与横截面联动", "单资产时间表示与跨资产注意力共同形成排序信号")
    assets = ["AAPL", "MSFT", "NVDA", "TSLA", "JPM"]
    for i, asset in enumerate(assets):
        y = 205 + i * 105
        rounded_box(draw, (95, y, 335, y + 70), asset, "#dceefb", font=FONT_BODY)
        rounded_box(draw, (440, y, 700, y + 70), "时间编码\nMamba-2 / Patch", "#e6f4ea", font=FONT_SMALL)
        arrow(draw, (335, y + 35), (440, y + 35))
        arrow(draw, (700, y + 35), (845, 445))
    rounded_box(draw, (845, 315, 1145, 575), "资产维编码\niTransformer\nCross-Asset Attention", "#fff4d6", font=FONT_BODY)
    rounded_box(draw, (1260, 350, 1510, 540), "相对强弱排序\nTop-N 选股信号", "#f3e8ff", font=FONT_BODY)
    arrow(draw, (1145, 445), (1260, 445))
    return save(img, "fig_4_4_cross_asset.png")


def fig_5_1() -> Path:
    img, draw = canvas("Walk-Forward 时序验证机制", "训练、验证、测试窗口沿时间轴逐轮向前推进")
    x0, y0 = 130, 230
    unit = 95
    for i in range(13):
        x = x0 + i * unit
        draw.rectangle((x, 705, x + unit - 8, 735), fill="#dbeafe")
        draw.text((x + 23, 748), f"T{i+1}", font=FONT_TINY, fill="#334155")
    rows = [
        (260, 0),
        (410, 2),
        (560, 4),
    ]
    for row_y, offset in rows:
        draw.text((80, row_y + 20), f"第{rows.index((row_y, offset))+1}轮", font=FONT_BODY, fill="#1f3d56")
        draw.rounded_rectangle((x0 + offset * unit, row_y, x0 + (offset + 5) * unit, row_y + 70), 12, fill="#bbf7d0", outline="#166534", width=3)
        draw.rounded_rectangle((x0 + (offset + 5) * unit, row_y, x0 + (offset + 7) * unit, row_y + 70), 12, fill="#fde68a", outline="#92400e", width=3)
        draw.rounded_rectangle((x0 + (offset + 7) * unit, row_y, x0 + (offset + 8) * unit, row_y + 70), 12, fill="#fecaca", outline="#991b1b", width=3)
        draw.text((x0 + offset * unit + 150, row_y + 20), "Train", font=FONT_BODY, fill="#166534")
        draw.text((x0 + (offset + 5) * unit + 42, row_y + 20), "Val", font=FONT_BODY, fill="#92400e")
        draw.text((x0 + (offset + 7) * unit + 20, row_y + 20), "Test", font=FONT_BODY, fill="#991b1b")
    return save(img, "fig_5_1_walk_forward.png")


def fig_5_2() -> Path:
    img, draw = canvas("Anti-Churn 集合差集调仓机制", "只对真正进入或退出组合的资产执行交易")
    rounded_box(draw, (120, 250, 520, 510), "昨日持仓集合 H(t-1)\n{AAPL, MSFT, NVDA, JPM}", "#dceefb", font=FONT_BODY)
    rounded_box(draw, (1080, 250, 1480, 510), "今日目标集合 H*(t)\n{MSFT, NVDA, TSLA, JPM}", "#e6f4ea", font=FONT_BODY)
    rounded_box(draw, (620, 185, 980, 330), "继续持有\n交集：{MSFT, NVDA, JPM}", "#fff4d6", font=FONT_BODY)
    rounded_box(draw, (620, 380, 980, 525), "卖出 / 买入\n差集：AAPL → TSLA", "#fde2e2", font=FONT_BODY)
    rounded_box(draw, (480, 645, 1120, 745), "减少边际排名抖动带来的无信息交易，降低换手率与成本侵蚀", "#ffffff", font=FONT_BODY)
    arrow(draw, (520, 380), (620, 260))
    arrow(draw, (1080, 380), (980, 260))
    arrow(draw, (520, 430), (620, 455))
    arrow(draw, (1080, 430), (980, 455))
    return save(img, "fig_5_2_anti_churn.png")


def fig_6_1() -> Path:
    img, draw = canvas("系统实现与运行链路", "FastAPI 后端与 Vue 3 前端共同支撑实验展示")
    rounded_box(draw, (100, 240, 380, 430), "PyTorch 模型\n模型注册\nCPU 降级", "#dceefb", font=FONT_BODY)
    rounded_box(draw, (510, 205, 790, 465), "FastAPI 服务\n/health\n/predictor/status\n/dashboard\n/trade", "#e6f4ea", font=FONT_BODY)
    rounded_box(draw, (920, 165, 1500, 270), "Dashboard：净值曲线、指标卡片、推荐列表", "#fff4d6", font=FONT_SMALL)
    rounded_box(draw, (920, 350, 1500, 455), "Analysis：K线走势、AI评分、技术指标", "#f3e8ff", font=FONT_SMALL)
    rounded_box(draw, (920, 535, 1500, 640), "Trade：仓位建议、手动下单、自动交易", "#fde2e2", font=FONT_SMALL)
    arrow(draw, (380, 335), (510, 335))
    arrow(draw, (790, 285), (920, 220))
    arrow(draw, (790, 335), (920, 402))
    arrow(draw, (790, 390), (920, 588))
    rounded_box(draw, (320, 700, 1280, 780), "异常被转换为结构化响应，前端展示友好提示，保证答辩演示稳定", "#ffffff", font=FONT_BODY)
    return save(img, "fig_6_1_runtime_chain.png")


def fig_6_2() -> Path:
    img, draw = canvas("AlphaTransformer 版本实验结果对比", "v1、v2、v3 在收益、风险和执行效率上的综合变化")
    metrics = [
        ("年化收益率", [18.6, 24.7, 31.4], "%", 35),
        ("Sharpe", [0.28, 1.05, 1.65], "", 2.0),
        ("最大回撤", [15.2, 11.6, 9.8], "%", 18),
        ("换手率", [68.0, 36.5, 18.7], "%", 75),
        ("交易次数", [147, 74, 35], "次", 160),
    ]
    colors = ["#93c5fd", "#86efac", "#fbbf24"]
    labels = ["v1", "v2", "v3"]
    chart_x, chart_y = 130, 220
    group_w = 270
    base_y = 710
    for i, (name, vals, unit, max_v) in enumerate(metrics):
        x = chart_x + i * group_w
        draw.text((x + 35, 760), name, font=FONT_SMALL, fill="#334155")
        for j, val in enumerate(vals):
            h = int((val / max_v) * 420)
            bx = x + 30 + j * 58
            draw.rectangle((bx, base_y - h, bx + 42, base_y), fill=colors[j], outline="#334155")
            draw.text((bx - 6, base_y - h - 34), f"{val:g}{unit}", font=FONT_TINY, fill="#334155")
    for j, label in enumerate(labels):
        draw.rectangle((1220, 140 + j * 48, 1260, 170 + j * 48), fill=colors[j], outline="#334155")
        draw.text((1275, 138 + j * 48), label, font=FONT_BODY, fill="#334155")
    draw.line((95, base_y, 1485, base_y), fill="#475569", width=3)
    return save(img, "fig_6_2_version_metrics.png")


def generate_figures() -> dict[str, Path]:
    return {
        "图3-1": fig_3_1(),
        "图3-2": fig_3_2(),
        "图4-1": fig_4_1(),
        "图4-2": fig_4_2(),
        "图4-3": fig_4_3(),
        "图4-4": fig_4_4(),
        "图5-1": fig_5_1(),
        "图5-2": fig_5_2(),
        "图6-1": fig_6_1(),
        "图6-2": fig_6_2(),
    }


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(9)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "666666")


def move_after(anchor, element):
    anchor._p.addnext(element)
    return element


def add_paragraph_after(doc: Document, anchor, text: str, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=10, bold=False):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(font_size)
    run.bold = bold
    move_after(anchor, p._p)
    return p


def add_picture_after(doc: Document, anchor, image_path: Path, caption: str):
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_p.add_run()
    run.add_picture(str(image_path), width=Inches(5.8))
    move_after(anchor, pic_p._p)
    cap_p = add_paragraph_after(doc, pic_p, caption, WD_ALIGN_PARAGRAPH.CENTER, 10, False)
    return cap_p


def add_table_after(doc: Document, anchor, caption: str, headers: list[str], rows: list[list[str]]):
    cap_p = add_paragraph_after(doc, anchor, caption, WD_ALIGN_PARAGRAPH.CENTER, 10, False)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    set_table_borders(table)
    cap_p._p.addnext(table._tbl)
    return table


def find_anchor(doc: Document, needle: str):
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            return paragraph
    raise ValueError(f"未找到插入锚点：{needle}")


def insert_content(doc: Document, figures: dict[str, Path]):
    tables = {
        "表2-1": (
            "表2-1 面向金融时序预测的模型方法对比",
            ["方法", "核心思想", "优势", "局限", "本文用途"],
            [
                ["标准 Transformer", "全局自注意力建模时序依赖", "表达能力强，可并行训练", "长序列成本高，易放大噪声", "作为早期基线"],
                ["PatchTST", "将时间序列划分为局部补丁", "降低序列长度，保留局部模式", "对市场状态切换感知不足", "v2 时间压缩思路"],
                ["iTransformer", "反转变量/资产维建模", "增强横截面联动表达", "需配合时间编码使用", "资产维关系建模"],
                ["Mamba", "选择性状态空间递推", "线性复杂度，适合长序列", "解释性依赖状态设计", "v3 时间记忆模块"],
                ["Regime-Aware", "识别市场状态并动态调整表示", "适应非平稳和波动切换", "状态划分需与任务匹配", "v3 状态感知模块"],
            ],
        ),
        "表3-1": (
            "表3-1 系统模块职责与代码对应关系",
            ["模块", "主要职责", "代码位置", "输出结果"],
            [
                ["数据预处理", "行情清洗、特征衍生、滚动标准化", "data/、utils/", "训练样本与特征矩阵"],
                ["模型训练", "版本模型训练、验证与参数管理", "models/、trainer/", "模型权重与预测分数"],
                ["回测评估", "防泄露对齐、交易成本、指标统计", "evaluation/", "收益曲线与评估指标"],
                ["推理服务", "模型加载、预测接口、异常封装", "api/", "结构化 API 响应"],
                ["可视化展示", "仪表盘、分析页、交易页", "frontend/", "图表与交互界面"],
            ],
        ),
        "表4-1": (
            "表4-1 AlphaTransformer 不同版本演进对比",
            ["版本", "核心结构", "主要解决问题", "IC", "Sharpe", "最大回撤", "换手率"],
            [
                ["v1", "原生 Transformer", "初步建模时序与资产关系", "0.073", "0.28", "-15.2%", "68.0%"],
                ["v2", "PatchTST + 稀疏连接", "缓解长序列成本和噪声扩散", "0.078", "1.05", "-11.6%", "36.5%"],
                ["v3", "Regime-Aware + Mamba-2 + iTransformer", "状态感知、时间记忆与横截面联动", "0.082", "1.65", "-9.8%", "18.7%"],
            ],
        ),
        "表4-2": (
            "表4-2 CombinedQuantLoss 构成与作用",
            ["损失项", "约束目标", "金融含义", "预期效果"],
            [
                ["预测误差项", "约束预测值与未来收益偏差", "保证基础数值拟合能力", "避免信号完全失真"],
                ["排序损失项", "强化横截面相对强弱识别", "服务 Top-N 选股与组合构建", "提高 IC 与 RankIC"],
                ["换手惩罚项", "抑制边际排名频繁翻转", "减少无信息调仓", "降低换手率"],
                ["交易成本项", "训练阶段引入摩擦约束", "让模型感知手续费与滑点", "保留扣费后净收益"],
            ],
        ),
        "表5-1": (
            "表5-1 实验数据与样本构建参数",
            ["项目", "设置", "说明"],
            [
                ["股票池", "多资产股票池", "用于横截面排序建模"],
                ["历史窗口", "60 个交易日", "构成模型输入时间长度"],
                ["预测周期", "未来收益/持仓周期", "作为监督目标和回测信号"],
                ["特征类型", "OHLCV、动量、波动、技术指标、截面因子", "覆盖价量与市场状态信息"],
                ["标准化方式", "滚动窗口标准化", "仅依赖当前及历史信息"],
                ["切分方式", "训练集早于验证集，验证集早于测试集", "避免随机切分造成时序泄露"],
            ],
        ),
        "表5-2": (
            "表5-2 实验指标体系",
            ["评价层面", "指标", "作用"],
            [
                ["预测层面", "MAE、RMSE", "判断数值预测误差是否失控"],
                ["排序层面", "IC、RankIC", "衡量横截面相对强弱识别能力"],
                ["收益风险层面", "年化收益、Sharpe Ratio、最大回撤", "评估策略收益质量和风险暴露"],
                ["交易执行层面", "换手率、交易次数、成本侵蚀比例", "判断信号是否具备实际可执行性"],
            ],
        ),
        "表6-1": (
            "表6-1 系统主要 API 接口",
            ["接口类别", "代表接口", "功能说明"],
            [
                ["系统状态", "/api/v1/health", "健康检查与服务运行状态"],
                ["模型管理", "/api/v1/predictor/status", "查看模型加载状态和设备信息"],
                ["模型预测", "/api/v1/predictor/load、预测相关接口", "加载模型并输出预测信号"],
                ["Dashboard", "/api/v1/dashboard/full、/dashboard/predictions", "返回权益曲线、指标与推荐列表"],
                ["账户管理", "/api/v1/account", "返回资金、持仓和账户统计"],
                ["交易执行", "/api/v1/trade/order、/trade/history", "下单、平仓与历史成交查询"],
                ["自动交易", "/api/v1/auto/status、/auto/start、/auto/stop", "自动交易状态控制与信号查看"],
            ],
        ),
        "表6-2": (
            "表6-2 v3 消融实验结果",
            ["实验设置", "IC", "Sharpe", "最大回撤", "换手率", "结论"],
            [
                ["完整 v3", "0.082", "1.65", "-9.8%", "18.7%", "收益、风险和执行效率最均衡"],
                ["去除 Regime-Aware", "变化不大", "1.21", "-13.7%", "未单独报告", "异常阶段风险缓释能力下降"],
                ["去除 Mamba-2", "0.079", "1.33", "未单独报告", "未单独报告", "有效时间记忆提炼能力下降"],
                ["去除 iTransformer", "0.076", "未单独报告", "未单独报告", "未单独报告", "截面资产联动建模贡献明显"],
                ["去除交易约束与 Anti-Churn", "0.081", "0.94", "未单独报告", "52.4%", "预测仍可用，但交易可执行性明显恶化"],
            ],
        ),
    }

    # 第2章
    anchor = find_anchor(doc, "状态空间模型也为金融长序列建模提供了思路")
    add_table_after(doc, anchor, *tables["表2-1"])

    # 第3章
    anchor = find_anchor(doc, "系统采用数据层、模型层、服务层和展示层四层架构")
    last = add_picture_after(doc, anchor, figures["图3-1"], "图3-1 系统四层总体架构图")
    anchor = find_anchor(doc, "系统关键流程可概括为三步")
    add_picture_after(doc, anchor, figures["图3-2"], "图3-2 核心业务流程图")
    anchor = find_anchor(doc, "系统内部划分为数据预处理、模型训练、回测评估、推理服务和可视化展示五类模块")
    add_table_after(doc, anchor, *tables["表3-1"])

    # 第4章
    anchor = find_anchor(doc, "简单说一下 v3 的处理流程")
    last = add_picture_after(doc, anchor, figures["图4-1"], "图4-1 AlphaTransformer v3 总体架构图")
    add_table_after(doc, last, *tables["表4-1"])
    anchor = find_anchor(doc, "状态概率不直接决定最终交易")
    add_picture_after(doc, anchor, figures["图4-2"], "图4-2 Regime-Aware 状态感知与动态归一化机制")
    anchor = find_anchor(doc, "Mamba-2 对“保留关键信号、衰减陈旧噪声”")
    add_picture_after(doc, anchor, figures["图4-3"], "图4-3 Mamba-2 时间编码与记忆衰减示意图")
    anchor = find_anchor(doc, "股票收益往往受行业景气、风格偏好和资金流向共同驱动")
    add_picture_after(doc, anchor, figures["图4-4"], "图4-4 时间资产维编码与横截面联动示意图")
    anchor = find_anchor(doc, "预测误差、排序损失与换手惩罚的组合形式")
    add_table_after(doc, anchor, *tables["表4-2"])

    # 第5章
    anchor = find_anchor(doc, "不采用随机打乱，也不用常规 K 折交叉验证")
    add_table_after(doc, anchor, *tables["表5-1"])
    anchor = find_anchor(doc, "训练窗口、验证窗口和测试窗口沿时间轴逐轮向前推进")
    add_picture_after(doc, anchor, figures["图5-1"], "图5-1 Walk-Forward 时序验证机制")
    anchor = find_anchor(doc, "系统只对真正进入或退出组合的资产执行交易")
    add_picture_after(doc, anchor, figures["图5-2"], "图5-2 Anti-Churn 集合差集调仓机制")
    anchor = find_anchor(doc, "预测、排序、收益风险和交易执行四个层面评估模型")
    add_table_after(doc, anchor, *tables["表5-2"])

    # 第6章
    anchor = find_anchor(doc, "后端采用 FastAPI 与 PyTorch 组合实现")
    last = add_picture_after(doc, anchor, figures["图6-1"], "图6-1 系统实现与运行链路")
    add_table_after(doc, last, *tables["表6-1"])
    anchor = find_anchor(doc, "v3 则把年化收益率提高到 31.4%")
    add_picture_after(doc, anchor, figures["图6-2"], "图6-2 AlphaTransformer 版本实验结果对比")
    anchor = find_anchor(doc, "若移除交易惩罚与 Anti-Churn 相关约束")
    add_table_after(doc, anchor, *tables["表6-2"])


def apply_document_styles(doc: Document):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("图") or paragraph.text.startswith("表"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0, 0, 0)


def main():
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)

    figures = generate_figures()
    backup_path = DOCX_PATH.with_name(f"22.15.backup_{datetime.now():%Y%m%d_%H%M%S}.docx")
    shutil.copy2(DOCX_PATH, backup_path)

    doc = Document(DOCX_PATH)
    if any(p.text.strip().startswith("图3-1 系统四层总体架构图") for p in doc.paragraphs):
        raise RuntimeError("文档中已存在本脚本插入的图表题注。请先使用原始备份恢复文档后再运行。")
    insert_content(doc, figures)
    apply_document_styles(doc)
    doc.save(DOCX_PATH)

    print(f"备份文件: {backup_path}")
    print(f"图片目录: {FIG_DIR}")
    print(f"已插入图片: {len(figures)}")
    print("已插入表格: 8")
    print(f"已更新文档: {DOCX_PATH}")


if __name__ == "__main__":
    main()
