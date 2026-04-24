# -*- coding: utf-8 -*-
"""最终验证 - 检查所有改写是否成功应用"""

from docx import Document

INPUT = r"D:/transformer/thesis/基于transformer的股票预测系统_降AIGC版.docx"

doc = Document(INPUT)
full_text = " ".join(p.text for p in doc.paragraphs)

print("=" * 60)
print("最终验证报告")
print("=" * 60)

# 1. 检查旧文本是否已被替换
OLD_TEXTS = [
    "股票价格预测是金融科技和量化投资交叉处的典型难题",
    "基于这些问题，本文设计并实现面向量化选股场景的 Transformer 股票预测系统。",
    "AlphaTransformer 的迭代不是简单叠加新模块",
    "这套指标结构，让" + chr(0x201C) + "预测模型研究" + chr(0x201D),
]

print("\n[1] 旧文本检查（应全部NOT FOUND）:")
for t in OLD_TEXTS:
    found = t in full_text
    print(f"  {'STILL EXISTS!' if found else 'OK - removed'} {t[:40]}...")

# 2. 检查新文本是否已写入
NEW_TEXTS = [
    "说它难，不光是因为行情本身涨涨跌跌不好猜",
    "针对上面提到的问题，本文决定自己动手做一套",
    "这几个版本的迭代，不是往上面随便加几个新模块就完事了",
    "研究就不只是" + chr(0x201C) + "训练了一个预测模型" + chr(0x201D),
]

print("\n[2] 新文本检查（应全部FOUND）:")
for t in NEW_TEXTS:
    found = t in full_text
    print(f"  {'FOUND' if found else 'MISSING!'} {t[:40]}...")

# 3. 关键数据完整性
print("\n[3] 关键数据完整性:")
checks = {
    "v3 Sharpe 1.65": "1.65",
    "IC 0.082": "0.082",
    "最大回撤 -9.8%": "-9.8%",
    "年化收益 31.4%": "31.4%",
    "换手率 18.7%": "18.7%",
    "累计交易 35 次": "35",
    "v1 Sharpe 0.28": "0.28",
    "v2 Sharpe 1.05": "1.05",
}
for label, val in checks.items():
    found = val in full_text
    print(f"  {'OK' if found else 'MISS!'} {label} ({val})")

# 4. 专业术语保留
print("\n[4] 专业术语保留:")
terms = [
    "Transformer", "Regime-Aware", "Mamba-2", "iTransformer",
    "AlphaTransformer", "CombinedQuantLoss", "Anti-Churn",
    "Sharpe Ratio", "Walk-Forward", "IC"
]
for term in terms:
    found = term in full_text
    print(f"  {'OK' if found else 'MISS!'} {term}")

# 5. 统计
print(f"\n[5] 文档统计:")
print(f"  总段落数: {len(doc.paragraphs)}")
print(f"  总字符数: {len(full_text)}")

print("\n" + "=" * 60)
print("验证完成")
print("=" * 60)
