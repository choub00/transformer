# -*- coding: utf-8 -*-
"""第二轮最终验证"""

from docx import Document

INPUT = r"D:/transformer/thesis/基于transformer的股票预测系统_降AIGC版.docx"
doc = Document(INPUT)
text = " ".join(p.text for p in doc.paragraphs)

print("=" * 60)
print("第二轮改写 - 最终验证报告")
print("=" * 60)

# 1. 新文本是否写入
NEW_CHECKS = [
    ("量化投资领域的核心难题之一", "摘要新写法"),
    ("简单来说", "说白了->简单来说"),
    ("实验结果也比较明显", "比较直观->也比较明显"),
    ("从这些数据可以看出", "这几个数字能看出来->从这些数据可以看出"),
    ("实验记录表明", "从实验记录来看->实验记录表明"),
    ("始终用同一套逻辑", "闷头用->始终用"),
    ("各有分工，而不是混在一起", "各管各的一摊->各有分工"),
    ("这几块内容不是简单拼在一起", "这一个东西->这几块内容"),
    ("并非简单增加新模块", "AlphaTransformer改写（两处）"),
    ("v3 成功将换手率控制到了", "v3能->v3成功将"),
    ("例如训练和回测", "比如说->例如"),
    ("关于分层架构已做说明", "分层架构说完了->关于分层架构已做说明"),
    ("难以落实到代码层面", "落实不到代码上->难以落实到代码层面"),
]
print("\n[1] 新文本验证:")
all_ok = True
for kw, label in NEW_CHECKS:
    ok = kw in text
    print(f"  {'OK' if ok else 'MISS'} {label}: '{kw}'")
    if not ok:
        all_ok = False

# 2. 旧文本是否清除
OLD_CHECKS = [
    ("金融科技和量化投资交叉处", "旧摘要第1句"),
    ("老问题", "老问题口语"),
    ("说白了", "说白了口语"),
    ("比较直观", "比较直观口语"),
    ("从这几个数字能看出来", "从这几个数字能看出来"),
    ("从实验记录来看", "从实验记录来看"),
    ("闷头", "闷头口语"),
    ("各管各的一摊", "各管各的一摊"),
    ("这一个东西", "这一个东西"),
    ("不是往上面加几个新模块", "Alpha旧版本"),
    ("v3 因此把换手率压到了", "v3 因而下沉旧写法"),
    ("比如说", "比如说口语"),
    ("落实不到代码上", "落实不到代码上"),
]
print("\n[2] 旧文本清除:")
for kw, label in OLD_CHECKS:
    ok = kw not in text
    print(f"  {'OK' if ok else 'STILL'} {label}: '{kw}'")

# 3. 关键数据
print("\n[3] 关键数据完整性:")
checks = {
    "v3 Sharpe 1.65": "1.65",
    "v3 IC 0.082": "0.082",
    "v3 MaxDD -9.8%": "-9.8%",
    "v3 年化 31.4%": "31.4%",
    "v3 换手 18.7%": "18.7%",
    "v3 交易 35次": "35",
    "v2 Sharpe 1.05": "1.05",
    "v1 Sharpe 0.28": "0.28",
    "v1 IC 0.073": "0.073",
}
all_data_ok = True
for label, val in checks.items():
    ok = val in text
    print(f"  {'OK' if ok else 'MISS'} {label} ({val})")
    if not ok:
        all_data_ok = False

# 4. 专业术语
print("\n[4] 专业术语保留:")
terms = ["Transformer", "Regime-Aware", "Mamba-2", "iTransformer",
         "AlphaTransformer", "CombinedQuantLoss", "Anti-Churn",
         "Walk-Forward", "Sharpe Ratio", "IC"]
for term in terms:
    ok = term in text
    print(f"  {'OK' if ok else 'MISS'} {term}")

print("\n" + "=" * 60)
print("文档段落数:", len(doc.paragraphs))
print("文档总字符:", len(text))
print("=" * 60)
