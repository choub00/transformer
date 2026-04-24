# -*- coding: utf-8 -*-
from docx import Document

OUTPUT = r"D:/transformer/thesis/基于transformer的股票预测系统_润色版.docx"

doc = Document(OUTPUT)

# Words to check
check_words = ["因此", "本课题"]
key_values = ["1.65", "0.082", "-9.8%", "31.4%", "18.7%", "35"]

# Count occurrences
word_counts = {w: 0 for w in check_words}
key_present = {v: False for v in key_values}

def check_text(text):
    for w in check_words:
        word_counts[w] += text.count(w)
    for v in key_values:
        if v in text:
            key_present[v] = True

print("=" * 60)
print("VERIFICATION RESULTS")
print("=" * 60)

# Check paragraphs
print("\n[1] Checking paragraphs...")
for i, para in enumerate(doc.paragraphs):
    check_text(para.text)

# Check tables
print("[2] Checking tables...")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            check_text(cell.text)

# Report forbidden words
print("\n" + "=" * 60)
print("FORBIDDEN WORDS CHECK")
print("=" * 60)
for word, count in word_counts.items():
    status = "FAIL" if count > 0 else "PASS"
    print(f"  '{word}': {count} occurrences [{status}]")

# Report key values
print("\n" + "=" * 60)
print("KEY VALUES CHECK")
print("=" * 60)
for val, present in key_present.items():
    status = "FOUND" if present else "MISSING"
    print(f"  '{val}': [{status}]")

# Final result
print("\n" + "=" * 60)
print("FINAL RESULT")
print("=" * 60)
all_good = all(count == 0 for count in word_counts.values()) and all(key_present.values())
if all_good:
    print("  ALL CHECKS PASSED!")
else:
    print("  SOME CHECKS FAILED - see above for details")
